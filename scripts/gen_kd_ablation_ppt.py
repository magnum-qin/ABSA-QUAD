"""Generate the final KD ablation evidence deck.

1. One-shot official-test KD comparisons
2. Size-matched Russian pool-source comparison
3. Same-pool gen8k quantity control
4. Six evidence-backed conclusions

Same visual language as gen_kd_ppt.py (clean pastel-panel theme, English).
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FONT_BODY = "Segoe UI"
FONT_MONO = "Consolas"

C_BG        = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE     = RGBColor(0x0F, 0x17, 0x2A)
C_BODY      = RGBColor(0x33, 0x41, 0x55)
C_MUTED     = RGBColor(0x64, 0x74, 0x8B)
C_EYEBROW   = RGBColor(0x25, 0x63, 0xEB)
C_RED       = RGBColor(0xDC, 0x26, 0x26)
C_GREEN     = RGBColor(0x05, 0x96, 0x69)
C_AMBER     = RGBColor(0xB4, 0x5A, 0x09)

PANEL_BLUE   = (RGBColor(0xDE, 0xEA, 0xFB), RGBColor(0x1D, 0x4E, 0xD8), RGBColor(0x3B, 0x82, 0xF6))
PANEL_GREEN  = (RGBColor(0xD6, 0xF2, 0xE4), RGBColor(0x04, 0x78, 0x57), RGBColor(0x10, 0xB9, 0x81))
PANEL_SLATE  = (RGBColor(0xEE, 0xF2, 0xF7), RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x64, 0x74, 0x8B))
PANEL_AMBER  = (RGBColor(0xFE, 0xF3, 0xC7), RGBColor(0xB4, 0x5A, 0x09), RGBColor(0xF5, 0x9E, 0x0B))
PANEL_RED    = (RGBColor(0xFD, 0xE4, 0xE4), RGBColor(0xDC, 0x26, 0x26), RGBColor(0xEF, 0x44, 0x44))

C_HEADER_ROW = RGBColor(0xF1, 0xF5, 0xF9)
C_ROW_LINE   = RGBColor(0xE2, 0xE8, 0xF0)
C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
LAYOUT = prs.slide_layouts[6]


def add_slide():
    s = prs.slides.add_slide(LAYOUT)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C_BG
    return s


def rrect(slide, x, y, w, h, fill, line=None, radius=0.06, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _spacing(run, spc):
    rPr = run._r.get_or_add_rPr()
    rPr.set('spc', str(spc))


def tx(slide, text, x, y, w, h, size=18, bold=False, color=C_BODY,
       align=PP_ALIGN.LEFT, italic=False, font=FONT_BODY, spacing=None,
       anchor=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    if anchor:
        tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    if spacing:
        _spacing(r, spacing)
    return tb


def eyebrow(slide, text, x=0.55, y=0.42):
    tx(slide, text.upper(), x, y, 11.0, 0.35, size=13, bold=True,
       color=C_EYEBROW, spacing=220, font=FONT_BODY)


def title(slide, text, x=0.55, y=0.74, size=28, w=12.2):
    tx(slide, text, x, y, w, 0.75, size=size, bold=True, color=C_TITLE)


def data_table(slide, x, y, col_w, headers, rows, size=13, header_size=11.5,
               row_h=0.48, highlight_idx=None, highlight_color=None,
               first_col_bold=True, first_col_color=None, font=FONT_BODY,
               left_align_cols=(2,), cell_color_fn=None):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    total_w = sum(col_w)
    rrect(slide, x, y, total_w, 0.42, C_HEADER_ROW, radius=0.08)
    for h, cx, w in zip(headers, col_x, col_w):
        tx(slide, h, cx + 0.08, y + 0.04, w - 0.16, 0.34, size=header_size, bold=True,
           color=C_MUTED, align=PP_ALIGN.CENTER, spacing=30)
    yy = y + 0.42
    for i, row in enumerate(rows):
        bg = highlight_color if (highlight_idx == i and highlight_color) else C_WHITE
        rrect(slide, x, yy, total_w, row_h, bg, line=C_ROW_LINE, radius=0.03)
        for j, (cell, cx, w) in enumerate(zip(row, col_x, col_w)):
            bold0 = first_col_bold and j == 0
            col0 = (first_col_color or C_TITLE) if j == 0 else C_BODY
            if cell_color_fn:
                override = cell_color_fn(i, j, cell)
                if override:
                    col0 = override
                    bold0 = True
            al = PP_ALIGN.LEFT if j in left_align_cols else PP_ALIGN.CENTER
            tx(slide, str(cell), cx + 0.1, yy, w - 0.2, row_h, size=size, bold=bold0,
               color=col0, align=al, anchor=MSO_ANCHOR.MIDDLE, font=font)
        yy += row_h
    return yy


# ══════════════════════════════════════════════════════════════
#  SLIDE 1 — Cover
# ══════════════════════════════════════════════════════════════
s = add_slide()
eyebrow(s, "ABSA-QUAD  ·  Final KD Evidence", y=1.6)
tx(s, "What Actually Controls KD Gain?", 0.55, 2.1, 12.2, 1.0,
   size=44, bold=True, color=C_TITLE)
tx(s, "Teacher × Pool × Quality–Coverage Trade-off", 0.55, 3.05, 12.2, 0.65,
   size=26, color=C_EYEBROW, bold=True)
rect(s, 0.6, 3.95, 1.7, 0.05, C_EYEBROW)
tx(s, "TALLIP evidence closure · clean-inductive + official test · 2026-07-18",
   0.55, 4.15, 12.2, 0.5, size=16, color=C_MUTED, italic=True)


# ══════════════════════════════════════════════════════════════
#  SLIDE 2 — One-shot official-test KD comparisons
# ══════════════════════════════════════════════════════════════
s = add_slide()
eyebrow(s, "Official Test · Pre-Locked")
title(s, "Four Supported KD Gains; One Inconclusive")

headers = ["DATASET", "KD", "NO-KD", "Δ cF1", "95% CI", "DECISION"]
col_w = [2.1, 1.55, 1.55, 1.5, 3.4, 2.1]
rows = [
    ("eng_laptop", "0.3532", "0.2710", "+0.0822", "[+0.0629, +0.1021]", "Supported"),
    ("rus_restaurant", "0.5189", "0.4978", "+0.0211", "[+0.0005, +0.0431]", "Supported"),
    ("tat_restaurant", "0.4311", "0.4037", "+0.0274", "[+0.0071, +0.0482]", "Supported"),
    ("ukr_restaurant", "0.4608", "0.4215", "+0.0393", "[+0.0150, +0.0639]", "Supported"),
    ("zho_restaurant", "0.4660", "0.4798", "−0.0137", "[−0.0289, +0.0014]", "Inconclusive"),
]

def _adopted_color(i, j, cell):
    if j == 5:
        return PANEL_GREEN[1] if cell == "Supported" else C_AMBER
    return None

bottom = data_table(s, 0.52, 1.82, col_w, headers, rows, row_h=0.70, size=13,
                     cell_color_fn=_adopted_color)

tx(s, "10,000 paired bootstrap resamples · seed 20260713. Systems, predictions and thresholds were "
      "frozen before gold retrieval. None of the eight main systems establishes test SOTA.",
   0.55, bottom + 0.25, 12.2, 0.65, size=13, italic=True, color=C_MUTED)


# ══════════════════════════════════════════════════════════════
#  SLIDE 3 — Pool -> Silver -> Training pipeline
# ══════════════════════════════════════════════════════════════
s = add_slide()
eyebrow(s, "Ablation Design · Dev Evidence")
title(s, "Size-Matched Pool Source Comparison (rus_restaurant)")

headers2 = ["CONTROL / RESULT", "REAL4K", "TRANSLATED4K", "GENERATED4K"]
col_w2 = [4.6, 2.55, 2.55, 2.55]
rows2 = [
    ("Pool size (unlabeled text)", "4,000", "3,466", "4,000"),
    ("DeepSeek-labeled silver", "4,000", "3,466", "4,000"),
    ("Silver used (cap=3000)", "3,000", "3,000", "3,000"),
    ("+ gold (fixed)", "1,240", "1,240", "1,240"),
    ("Final training rows", "4,240", "4,240", "4,240"),
    ("5-seed ensemble cF1", "0.5586", "0.5761", "0.6162"),
]
data_table(s, 0.52, 2.0, col_w2, headers2, rows2, row_h=0.62, size=14.5, header_size=13)


# ══════════════════════════════════════════════════════════════
#  SLIDE 4 — Same-pool quantity control
# ══════════════════════════════════════════════════════════════
s = add_slide()
eyebrow(s, "Quantity Control · Dev Evidence")
title(s, "Same gen8k Silver: Reducing to 3,000 Does Not Help")

headers3 = ["PAIRED EVIDENCE", "FULL GEN8K (7,988)", "CAP 3,000"]
col_w3 = [2.6, 4.85, 4.85]
rows3 = [
    ("Seed 42", "0.5286", "0.5123"),
    ("Seed 123", "0.5648", "0.5584"),
    ("Paired mean", "0.5467", "0.5353"),
    ("Δ vs full", "—", "−1.14pp (gate failed)"),
]
data_table(s, 0.52, 2.0, col_w3, headers3, rows3, row_h=0.85, size=13, header_size=12)


# ══════════════════════════════════════════════════════════════
#  SLIDE 5 — Final synthesis
# ══════════════════════════════════════════════════════════════
s = add_slide()
eyebrow(s, "Final Synthesis")
title(s, "Six Conclusions Within the Evidence Boundary")

conclusions = [
    ("1", "KD gains are conditional", "Complementary, accurate coverage is a supported interpretation—not a universal law.", PANEL_BLUE),
    ("2", "DeepSeek is more robust overall", "Direct fidelity is 8/8 higher by point estimate; only 3/8 gaps are supported.", PANEL_AMBER),
    ("3", "Pool source has no universal ranking", "zho: real wins; rus: generated wins under size matching.", PANEL_GREEN),
    ("4", "rus gen4k is a batch-quality effect", "Same-pool reduction failed; the 4k advantage is not a generic cap effect.", PANEL_RED),
    ("5", "Best silver balances quality and coverage", "More is not automatically better; no universal cap is supported.", PANEL_SLATE),
    ("6", "Official test supports 4/5 contrasts", "zho is inconclusive; no main system exceeds the leaderboard best.", PANEL_BLUE),
]
for i, (num, heading, body, palette) in enumerate(conclusions):
    y = 1.48 + i * 0.91
    rrect(s, 0.62, y, 12.05, 0.73, palette[0], radius=0.05)
    rrect(s, 0.82, y + 0.10, 0.54, 0.52, palette[2], radius=0.10)
    tx(s, num, 0.82, y + 0.10, 0.54, 0.52, size=17, bold=True, color=C_WHITE,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tx(s, heading, 1.58, y + 0.06, 4.35, 0.56, size=15.2, bold=True,
       color=palette[1], anchor=MSO_ANCHOR.MIDDLE)
    tx(s, body, 6.02, y + 0.06, 6.20, 0.56, size=12.8, color=C_BODY,
       anchor=MSO_ANCHOR.MIDDLE)


out = "/home/g24tka12/ABSA-QUAD/docs/KD_ablation_final_20260713.pptx"
prs.save(out)
print(f"Saved -> {out}  ({len(prs.slides)} slides)")
