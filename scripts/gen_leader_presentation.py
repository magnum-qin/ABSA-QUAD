#!/usr/bin/env python3
"""Generate a presentation-first LEADER-DimASQP paper talk.

The deck is intentionally designed for projection: large type, one message per
slide, and an explicit task -> method -> evidence -> limitations narrative.
"""

from pathlib import Path
import subprocess
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "LEADER_DimASQP_Presentation_revised.pptx"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
FONT_LATIN = "Aptos"

NAVY = "173B5E"
INK = "172033"
MUTED = "5D6B7E"
TEAL = "14A39A"
BLUE = "2F6FDB"
GREEN = "2E9B67"
ORANGE = "ED9B40"
RED = "D85B66"
PURPLE = "7A63C5"
WHITE = "FFFFFF"
BG = "F7F9FC"
PALE_BLUE = "EAF1FC"
PALE_TEAL = "E6F6F3"
PALE_GREEN = "E8F5ED"
PALE_ORANGE = "FFF2DE"
PALE_RED = "FCEAEC"
PALE_PURPLE = "F0ECFA"
GRID = "D8E0EA"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_rect(slide, x, y, w, h, fill=WHITE, line=GRID, radius=True, lw=1.0):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line)
    sh.line.width = Pt(lw)
    if radius:
        try:
            sh.adjustments[0] = 0.10
        except Exception:
            pass
    return sh


def add_line(slide, x1, y1, x2, y2, color=GRID, width=1.4, dash=None):
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = dash
    return ln


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=20,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font=FONT,
    margin=0.04,
    italic=False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    p.line_spacing = 1.05
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs, x, y, w, h, size=20, color=INK, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.05
    for spec in runs:
        r = p.add_run()
        r.text = spec[0]
        r.font.name = spec[4] if len(spec) > 4 else FONT
        r.font.size = Pt(spec[1] if len(spec) > 1 else size)
        r.font.bold = spec[2] if len(spec) > 2 else False
        r.font.color.rgb = rgb(spec[3] if len(spec) > 3 else color)
    return box


def add_bullets(slide, items, x, y, w, h, size=19, color=INK, bullet_color=None, gap=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        r0 = p.add_run()
        r0.text = "●  "
        r0.font.name = FONT
        r0.font.size = Pt(max(size - 2, 14))
        r0.font.color.rgb = rgb(bullet_color or TEAL)
        r = p.add_run()
        r.text = item
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(color)
    return box


def add_title(slide, kicker, title, number=None):
    add_text(slide, kicker.upper(), 0.65, 0.30, 10.7, 0.28, 11.5, BLUE, True, font=FONT_LATIN)
    add_text(slide, title, 0.65, 0.63, 12.0, 0.55, 28, INK, True)
    add_line(slide, 0.65, 1.25, 12.68, 1.25, GRID, 1.2)
    if number is not None:
        add_text(slide, f"{number:02d}", 12.25, 0.28, 0.42, 0.25, 10.5, MUTED, True, PP_ALIGN.RIGHT, font=FONT_LATIN)


def add_footer(slide, text="LEADER-DimASQP · paper presentation"):
    add_text(slide, text, 0.68, 7.18, 10.8, 0.18, 9.5, MUTED, font=FONT_LATIN)


def add_pill(slide, text, x, y, w, fill=PALE_BLUE, color=BLUE, size=13):
    add_rect(slide, x, y, w, 0.34, fill, fill, True, 0)
    add_text(slide, text, x + 0.05, y + 0.02, w - 0.1, 0.25, size, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_card(slide, title, body, x, y, w, h, accent=TEAL, fill=WHITE, title_size=19, body_size=17):
    add_rect(slide, x, y, w, h, fill, GRID, True, 1)
    add_rect(slide, x, y, 0.08, h, accent, accent, False, 0)
    add_text(slide, title, x + 0.24, y + 0.18, w - 0.42, 0.34, title_size, accent, True)
    add_text(slide, body, x + 0.24, y + 0.62, w - 0.42, h - 0.76, body_size, INK)


def add_arrow(slide, x, y, w=0.42, h=0.28, color=BLUE):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    return sh


def pipeline_png():
    tmp = Path(tempfile.gettempdir()) / "leader_pipeline_presentation"
    png = tmp.with_suffix(".png")
    subprocess.run(
        ["pdftoppm", "-png", "-singlefile", "-r", "180", str(ROOT / "docs" / "leader_pipeline.pdf"), str(tmp)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return png


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1. Title
    s = prs.slides.add_slide(blank)
    set_bg(s, NAVY)
    add_text(s, "SEMEVAL-2026 TASK 3 · TRACK A SUBTASK 3", 0.72, 0.55, 7.2, 0.3, 12, "9BC2E6", True, font=FONT_LATIN)
    add_text(s, "LEADER-DimASQP", 0.72, 1.48, 11.8, 0.75, 43, WHITE, True, font=FONT_LATIN)
    add_text(s, "面向维度 ASQP 的语言特定编码器与数据感知集成系统", 0.72, 2.38, 11.5, 0.55, 26, "DCE8F5", True)
    add_text(s, "Language-Specific Encoder and Data-Aware Ensemble Refinement", 0.72, 3.05, 10.8, 0.36, 17, "AFC5DC", italic=True, font=FONT_LATIN)
    add_line(s, 0.72, 3.63, 3.05, 3.63, TEAL, 4)
    add_text(s, "核心结论", 0.72, 4.22, 1.5, 0.3, 14, "70D7CD", True)
    add_text(s, "没有一个统一配方适合所有语言；稳定的集成与经验证的编码器选择决定最终表现。", 0.72, 4.65, 10.9, 0.82, 23, WHITE, True)
    add_text(s, "汇报人：秦洪振 · 铃木美美 · 福本文代 · 周仁杰", 0.72, 6.35, 8.5, 0.3, 15.5, "DCE8F5")
    add_text(s, "杭州电子科技大学  ·  山梨大学", 0.72, 6.72, 7.0, 0.26, 13, "AFC5DC")

    # 2. Roadmap
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Talk roadmap", "用一条问题链讲清这篇论文", 2)
    steps = [
        ("01", "任务是什么？", "抽取结构 + 连续情绪回归", BLUE),
        ("02", "为什么困难？", "严格匹配、数据不平衡、回归收缩", PURPLE),
        ("03", "系统怎么做？", "Rethink、集合匹配、语言特定配置", TEAL),
        ("04", "证据是否充分？", "基线、逐数据集结果与消融", ORANGE),
        ("05", "还能改什么？", "失败经验、误差与局限", RED),
    ]
    x0, gap, cw = 0.68, 0.18, 2.28
    for i, (num, title, body, c) in enumerate(steps):
        x = x0 + i * (cw + gap)
        add_rect(s, x, 2.05, cw, 3.55, WHITE, GRID, True, 1)
        add_text(s, num, x + 0.18, 2.27, 0.55, 0.35, 17, c, True, font=FONT_LATIN)
        add_rect(s, x + 0.18, 2.86, 0.52, 0.08, c, c, False, 0)
        add_text(s, title, x + 0.18, 3.20, cw - 0.36, 0.6, 21, INK, True)
        add_text(s, body, x + 0.18, 4.12, cw - 0.36, 1.0, 17, MUTED)
    add_text(s, "叙事主线：从任务约束出发，解释设计选择，再用实验验证，而不是按论文目录逐页复述。", 1.1, 6.28, 11.1, 0.46, 19, NAVY, True, PP_ALIGN.CENTER)
    add_footer(s)

    # 3. Task and metric
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Task & metric", "DimASQP：先抽对结构，才有资格评价情绪距离", 3)
    add_text(s, "输入句子（示意）", 0.75, 1.55, 3.2, 0.3, 15, BLUE, True)
    add_rect(s, 0.72, 1.94, 5.15, 1.00, PALE_BLUE, "C7D8F2", True, 1)
    add_text(s, "“The food was excellent, but the service was cold.”", 0.98, 2.18, 4.65, 0.52, 21, INK, True, valign=MSO_ANCHOR.MIDDLE, font=FONT_LATIN)
    add_text(s, "输出集合：每条记录写成 (a, c, o, (V, A))", 0.75, 3.25, 5.0, 0.38, 17, MUTED)
    add_card(s, "Quad 1", "a: food\nc: FOOD#QUALITY\no: excellent  ·  V–A: 8.2, 6.4", 0.72, 3.76, 2.48, 1.63, BLUE, WHITE, 18, 14)
    add_card(s, "Quad 2", "a: service\nc: SERVICE#GENERAL\no: cold  ·  V–A: 2.8, 4.1", 3.38, 3.76, 2.48, 1.63, PURPLE, WHITE, 18, 14)
    add_rect(s, 6.35, 1.55, 6.22, 4.76, WHITE, GRID, True, 1)
    add_text(s, "cF1 的两道门", 6.70, 1.87, 5.4, 0.4, 22, NAVY, True)
    add_pill(s, "门 1 · categorical exact match", 6.70, 2.55, 2.70, PALE_RED, RED, 13)
    add_text(s, "Aspect + Category + Opinion 必须匹配", 6.72, 3.02, 5.25, 0.48, 20, INK, True)
    add_arrow(s, 8.80, 3.68, 0.55, 0.34, RED)
    add_pill(s, "门 2 · V–A distance discount", 9.58, 2.55, 2.60, PALE_TEAL, TEAL, 13)
    add_text(s, "通过门 1 后，再按 V–A 欧氏距离折扣", 7.20, 4.24, 4.55, 0.60, 20, INK, True, PP_ALIGN.CENTER)
    add_text(s, "cTP = 1 − distance / Dmax", 7.26, 5.22, 4.45, 0.46, 22, TEAL, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "因此：边界或类别错误会直接失去该条匹配；V–A 误差则连续扣分。", 0.88, 6.31, 11.65, 0.42, 19, NAVY, True, PP_ALIGN.CENTER)
    add_footer(s, "Metric definition: paper §3 · example values are illustrative")

    # 4. Data and challenges
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Data & challenges", "八个语言-领域数据集，规模与难点并不相同", 4)
    data = [
        ("中餐", 6050, TEAL), ("英笔记本", 4076, BLUE), ("中笔记本", 3490, BLUE),
        ("英餐", 2284, PURPLE), ("日酒店", 1600, ORANGE), ("俄餐", 1240, RED),
        ("鞑餐", 1240, RED), ("乌餐", 1240, RED),
    ]
    add_text(s, "训练句子数", 0.72, 1.53, 2.4, 0.3, 16, MUTED, True)
    maxv = max(v for _, v, _ in data)
    for i, (name, value, c) in enumerate(data):
        y = 1.93 + i * 0.54
        add_text(s, name, 0.72, y + 0.02, 0.85, 0.25, 14.5, INK, True, PP_ALIGN.RIGHT)
        add_rect(s, 1.70, y, 3.72, 0.31, "E8EDF3", "E8EDF3", True, 0)
        bw = 3.72 * value / maxv
        add_rect(s, 1.70, y, bw, 0.31, c, c, True, 0)
        add_text(s, f"{value:,}", 5.50, y - 0.01, 0.65, 0.26, 13.5, MUTED, True, font=FONT_LATIN)
    add_text(s, "最大 / 最小 ≈ 4.9×", 1.68, 6.38, 3.35, 0.35, 17, RED, True, PP_ALIGN.CENTER)
    add_card(s, "挑战 1 · 严格结构匹配", "Aspect / Opinion 边界只偏 1–2 个词，也可能无法形成 categorical match。", 6.30, 1.58, 6.15, 1.40, BLUE, WHITE, 19, 17.5)
    add_card(s, "挑战 2 · 低资源与跨语差异", "训练规模从 1,240 到 6,050；分词、形态与领域标签分布各不相同。", 6.30, 3.18, 6.15, 1.40, PURPLE, WHITE, 19, 17.5)
    add_card(s, "挑战 3 · V–A 中心化", "点回归容易收缩到中性区域，极端情绪的区分度不足。", 6.30, 4.78, 6.15, 1.40, ORANGE, WHITE, 19, 17.5)
    add_footer(s, "Dataset sizes: paper Table 1")

    # 5. System overview
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "System overview", "LEADER-DimASQP = 稳定底座 + 按数据集选择", 5)
    pimg = pipeline_png()
    s.shapes.add_picture(str(pimg), Inches(0.66), Inches(1.48), width=Inches(8.15))
    add_card(s, "A. 架构底座", "先做 draft，再把结构与 V–A 软预测反馈给 token 表征；候选四元组用 Hungarian set matching 训练。", 9.05, 1.58, 3.62, 1.72, BLUE, WHITE, 18, 16.5)
    add_card(s, "B. 数据与迁移", "回译、paraphrase、COMET 过滤与跨语热启动都只在验证有效的数据集上保留。", 9.05, 3.49, 3.62, 1.48, TEAL, WHITE, 18, 16.5)
    add_card(s, "C. 选择与推理", "按 dev cF1 选择编码器/阈值；五种子投票与 top-k 解码提升稳定性。", 9.05, 5.17, 3.62, 1.48, ORANGE, WHITE, 18, 16.5)
    add_footer(s, "Architecture figure: paper Figure 1")

    # 6. Architecture detail
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Architecture", "两项结构设计：Rethink 反馈与集合级匹配", 6)
    add_rect(s, 0.70, 1.55, 5.95, 4.85, PALE_BLUE, "C7D8F2", True, 1)
    add_text(s, "1 · Lightweight Rethink", 1.02, 1.85, 5.2, 0.4, 22, BLUE, True, font=FONT_LATIN)
    labels = [("Token states", 1.00), ("Draft heads", 2.45), ("Feedback", 3.90), ("Refined heads", 5.25)]
    for i, (lab, x) in enumerate(labels):
        add_rect(s, x, 2.62, 1.18, 0.76, WHITE, "AFC5DC", True, 1)
        add_text(s, lab, x + 0.05, 2.78, 1.08, 0.38, 13.5 if lab == "Feedback" else 15.5, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, font=FONT_LATIN)
        if i < len(labels) - 1:
            add_arrow(s, x + 1.22, 2.86, 0.24, 0.20, BLUE)
    add_bullets(s, ["BIO、类别、pair 与 V–A 草稿预测被嵌入", "与原 token 表征融合后，由第二阶段预测头细化"], 1.02, 3.86, 5.18, 1.55, 17.5, bullet_color=BLUE, gap=8)
    add_rect(s, 6.88, 1.55, 5.78, 4.85, PALE_ORANGE, "F1D6AC", True, 1)
    add_text(s, "2 · Hungarian set matching", 7.20, 1.85, 5.1, 0.4, 22, ORANGE, True, font=FONT_LATIN)
    add_rect(s, 7.20, 2.54, 1.48, 0.80, WHITE, "E4C28E", True, 1)
    add_text(s, "Pred set", 7.26, 2.75, 1.36, 0.30, 16, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, font=FONT_LATIN)
    add_rect(s, 10.80, 2.54, 1.48, 0.80, WHITE, "E4C28E", True, 1)
    add_text(s, "Gold\nset", 10.86, 2.68, 1.36, 0.45, 17, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, font=FONT_LATIN)
    add_arrow(s, 8.96, 2.78, 1.48, 0.30, ORANGE)
    add_text(s, "Cost = class + span + Huber(V–A)", 7.38, 3.70, 4.75, 0.40, 17, INK, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_bullets(s, ["在集合层面寻找一对一最优指派", "减少候选顺序与重复配对带来的训练不稳定"], 7.23, 4.35, 5.0, 1.20, 17.5, bullet_color=ORANGE, gap=8)
    add_text(s, "训练稳定器：FGM ε=0.1  ·  EMA 0.999  ·  top-3 checkpoints", 1.06, 6.65, 11.18, 0.35, 18.5, NAVY, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_footer(s, "Method details: paper §4.1–4.3")

    # 7. Dataset-specific selection
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Per-dataset selection", "最终系统不是一个模型，而是八个经验证的配置", 7)
    cfgs = [
        ("英餐", "DeBERTa-v3-large", "强抽取底座", BLUE, "0.6949"),
        ("英笔记本", "BT dual-layer ensemble", "保留历史最优", PURPLE, "0.3304"),
        ("日酒店", "Japanese DeBERTa", "MeCab 裁剪；no para", ORANGE, "0.4500"),
        ("俄餐", "XLM-R-large", "优于 ruRoberta", RED, "0.5267"),
        ("鞑餐", "XLM-R-large", "不保留 paraphrase", RED, "0.4821"),
        ("乌餐", "mDeBERTa-base", "BT + paraphrase", PURPLE, "0.5068"),
        ("中餐", "validated source model", "五种子集成", TEAL, "0.5669"),
        ("中笔记本", "Chinese RoBERTa-large", "top-k = 0.10", TEAL, "0.4487"),
    ]
    for i, (ds, enc, detail, c, score) in enumerate(cfgs):
        col, row = i % 2, i // 2
        x, y = 0.72 + col * 6.08, 1.58 + row * 1.30
        add_rect(s, x, y, 5.82, 1.10, WHITE, GRID, True, 1)
        add_rect(s, x, y, 0.10, 1.10, c, c, False, 0)
        add_text(s, ds, x + 0.24, y + 0.16, 1.05, 0.30, 18, INK, True)
        add_text(s, enc, x + 1.34, y + 0.14, 3.25, 0.30, 17, c, True, font=FONT_LATIN)
        add_text(s, detail, x + 1.34, y + 0.56, 3.50, 0.25, 15.5, MUTED)
        add_text(s, score, x + 4.48, y + 0.32, 1.08, 0.32, 17, NAVY, True, PP_ALIGN.RIGHT, font=FONT_LATIN)
    add_rect(s, 1.15, 6.88, 11.05, 0.08, TEAL, TEAL, False, 0)
    add_text(s, "原则：语言专用模型不一定更好；所有选择都必须由 dev cF1 验证。", 1.15, 6.45, 11.0, 0.35, 19, NAVY, True, PP_ALIGN.CENTER)
    add_footer(s, "Best validated per-dataset configuration · hidden-test labels not used")

    # 8. Augmentation nuance
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Data-aware augmentation", "增强不是越多越好：质量阈值与模型能力存在交互", 8)
    add_text(s, "日本语 paraphrase + COMET（mDeBERTa 阶段）", 0.75, 1.55, 5.65, 0.36, 18, BLUE, True)
    chart_x, chart_y, chart_w, chart_h = 0.92, 2.12, 5.45, 3.35
    add_line(s, chart_x, chart_y + chart_h, chart_x + chart_w, chart_y + chart_h, MUTED, 1.2)
    add_line(s, chart_x, chart_y, chart_x, chart_y + chart_h, MUTED, 1.2)
    vals = [("0.92", 0.3317), ("0.95", 0.3523), ("0.96", 0.3565), ("0.965", 0.3418)]
    vmin, vmax = 0.325, 0.36
    pts = []
    for i, (lab, val) in enumerate(vals):
        x = chart_x + 0.72 + i * 1.40
        y = chart_y + chart_h - 0.42 - (val - vmin) / (vmax - vmin) * 2.50
        pts.append((x, y))
        add_rect(s, x - 0.08, y - 0.08, 0.16, 0.16, TEAL if lab != "0.965" else RED, TEAL if lab != "0.965" else RED, True, 0)
        add_text(s, f"{val:.4f}", x - 0.42, y - 0.45, 0.84, 0.28, 14, INK, True, PP_ALIGN.CENTER, font=FONT_LATIN)
        add_text(s, lab, x - 0.40, chart_y + chart_h + 0.12, 0.80, 0.24, 13.5, MUTED, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        add_line(s, x1, y1, x2, y2, TEAL, 2.2)
    add_text(s, "COMET threshold", 2.48, 5.95, 2.3, 0.25, 14, MUTED, font=FONT_LATIN)
    add_pill(s, "sweet spot", 3.98, 1.90, 1.08, PALE_TEAL, TEAL, 12)
    add_card(s, "结论 1 · 倒 U 型", "0.96 是该阶段的最佳阈值；提高到 0.965 后，样本量下降并出现欠拟合。", 7.05, 1.62, 5.55, 1.48, TEAL, WHITE, 19, 17)
    add_card(s, "结论 2 · 不能跨模型照搬", "换成 Japanese DeBERTa + 五种子集成后，no paraphrase 反而更好：0.4500 vs 0.4466。", 7.05, 3.34, 5.55, 1.48, ORANGE, WHITE, 19, 17)
    add_card(s, "结论 3 · 迁移也需验证", "跨语热启动在早期低资源设置有效；最终俄/鞑使用 XLM-R-large，乌克兰保留 BT + paraphrase。", 7.05, 5.06, 5.55, 1.48, PURPLE, WHITE, 19, 17)
    add_footer(s, "COMET ablation: paper §6.4 · final Japanese recipe uses no paraphrase")

    # 9. Main mean results
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Main results", "平均 dev cF1 = 0.5008，显著超过官方基线", 9)
    systems = [
        ("LEADER-DimASQP", 0.5008, TEAL),
        ("Shared-task average", 0.4056, MUTED),
        ("Llama-3.3 70B", 0.3862, BLUE),
        ("GPT-OSS 120B", 0.3727, PURPLE),
        ("Kimi K2 0-shot", 0.2542, ORANGE),
        ("GPT-5 mini 0-shot", 0.2254, RED),
    ]
    for i, (name, val, c) in enumerate(systems):
        y = 1.72 + i * 0.76
        add_text(s, name, 0.72, y + 0.05, 2.55, 0.26, 15.5 if i else 17, INK, i == 0, font=FONT_LATIN)
        add_rect(s, 3.24, y, 6.45, 0.40, "E8EDF3", "E8EDF3", True, 0)
        add_rect(s, 3.24, y, 6.45 * val / 0.55, 0.40, c, c, True, 0)
        add_text(s, f"{val:.4f}", 9.72, y - 0.01, 0.95, 0.30, 15.5, c, True, PP_ALIGN.RIGHT, font=FONT_LATIN)
    add_rect(s, 10.76, 1.70, 1.88, 1.42, PALE_TEAL, TEAL, True, 1)
    add_text(s, "+11.46", 10.88, 1.96, 1.64, 0.42, 23, TEAL, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "pp vs Llama", 10.88, 2.50, 1.64, 0.25, 13, MUTED, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_rect(s, 10.76, 3.42, 1.88, 1.42, PALE_BLUE, BLUE, True, 1)
    add_text(s, "+9.52", 10.88, 3.68, 1.64, 0.42, 25, BLUE, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "pp vs task avg.", 10.88, 4.22, 1.64, 0.25, 13, MUTED, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_rect(s, 10.76, 5.14, 1.88, 1.42, PALE_ORANGE, ORANGE, True, 1)
    add_text(s, "8 / 8", 10.88, 5.40, 1.64, 0.42, 25, ORANGE, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "beats official\nbest baseline", 10.88, 5.91, 1.64, 0.42, 12.5, MUTED, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_footer(s, "Development-set comparison · official baselines from DimABSA paper")

    # 10. Dataset results and external comparison
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Per-dataset evidence", "优势并不均匀：三项刷新已报告 top，四项仍有差距", 10)
    rows = [
        ("Eng-R", 0.6949, "strong", GREEN),
        ("Eng-L", 0.3304, "below reported best", RED),
        ("Jpn-H", 0.4500, "reported top ↑", GREEN),
        ("Rus-R", 0.5267, "below reported best", RED),
        ("Tat-R", 0.4821, "reported top ↑", GREEN),
        ("Ukr-R", 0.5068, "below reported best", RED),
        ("Zho-R", 0.5669, "reported top ↑", GREEN),
        ("Zho-L", 0.4487, "below reported best", RED),
    ]
    add_rect(s, 0.74, 1.55, 7.50, 5.25, WHITE, GRID, True, 1)
    for i, (name, score, status, c) in enumerate(rows):
        y = 1.78 + i * 0.60
        if i:
            add_line(s, 0.98, y - 0.12, 7.98, y - 0.12, "E9EDF2", 0.8)
        add_text(s, name, 1.02, y, 1.10, 0.25, 16, INK, True, font=FONT_LATIN)
        add_text(s, f"{score:.4f}", 2.45, y, 1.05, 0.25, 17, NAVY, True, font=FONT_LATIN)
        add_pill(s, status, 4.15, y - 0.03, 2.75, PALE_GREEN if c == GREEN else PALE_RED, c, 12.5)
    add_card(s, "论文可以自信地说", "• 超过全部官方开发集基线\n• 平均 cF1 提升显著\n• Jpn / Tat / Zho-R 超过已报告 top", 8.55, 1.68, 4.05, 2.06, TEAL, WHITE, 19, 17.5)
    add_card(s, "论文不应说", "• “所有数据集全面领先”\n• “绝对统治力”\n• 将 dev 结果写成 hidden-test 排名", 8.55, 4.05, 4.05, 2.06, RED, WHITE, 19, 17.5)
    add_text(s, "学术表达要区分：official baseline / reported system / hidden test。", 8.58, 6.45, 4.00, 0.32, 15.5, MUTED, True, PP_ALIGN.CENTER)
    add_footer(s, "External scores are reference comparisons, not hidden-test claims")

    # 11. Ablations
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Ablation studies", "最稳健的最终贡献来自集成；其他收益依赖模型规模", 11)
    add_rect(s, 0.72, 1.58, 4.15, 4.93, PALE_TEAL, "BEE5DE", True, 1)
    add_text(s, "5-seed ensemble", 1.02, 1.92, 3.55, 0.42, 23, TEAL, True, font=FONT_LATIN)
    add_text(s, "+4.76 pp", 1.02, 2.62, 3.55, 0.62, 35, TEAL, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "相对五种子单模均值", 1.02, 3.33, 3.55, 0.32, 17, MUTED, True, PP_ALIGN.CENTER)
    add_bullets(s, ["8 个数据集全部提升", "单项增益 +1.43 ～ +7.74 pp", "边界投票 + V–A 均值降低方差"], 1.06, 4.05, 3.48, 1.60, 17, bullet_color=TEAL, gap=6)
    add_rect(s, 5.12, 1.58, 3.55, 4.93, PALE_BLUE, "C7D8F2", True, 1)
    add_text(s, "Rethink", 5.42, 1.92, 2.90, 0.42, 23, BLUE, True, font=FONT_LATIN)
    add_text(s, "+11.29", 5.40, 2.62, 2.94, 0.52, 29, BLUE, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "pp · early base", 5.40, 3.18, 2.94, 0.26, 14, MUTED, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_line(s, 5.72, 3.74, 8.06, 3.74, "AFC5DC", 1)
    add_text(s, "+0.36", 5.40, 4.07, 2.94, 0.52, 29, BLUE, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "pp · large + ensemble", 5.40, 4.63, 2.94, 0.28, 14, MUTED, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "作用随 encoder/ensemble 变强而衰减", 5.45, 5.47, 2.84, 0.54, 16.5, NAVY, True, PP_ALIGN.CENTER)
    add_rect(s, 8.92, 1.58, 3.70, 4.93, PALE_ORANGE, "F1D6AC", True, 1)
    add_text(s, "Native encoder", 9.22, 1.92, 3.10, 0.42, 23, ORANGE, True, font=FONT_LATIN)
    enc = [("Japanese", "+5.04", GREEN), ("Chinese-L", "+5.83", GREEN), ("Russian", "−1.15", RED)]
    for i, (lab, val, c) in enumerate(enc):
        y = 2.75 + i * 0.82
        add_text(s, lab, 9.26, y, 1.55, 0.30, 16.5, INK, True, font=FONT_LATIN)
        add_text(s, val, 11.02, y, 1.08, 0.30, 16.5, c, True, PP_ALIGN.RIGHT, font=FONT_LATIN)
    add_text(s, "结论：language-specific ≠ automatically better", 9.28, 5.50, 3.00, 0.52, 16, NAVY, True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_text(s, "论文主线应强调“验证驱动”，而不是把所有组件都包装成普适增益。", 1.08, 6.73, 11.18, 0.32, 18.5, NAVY, True, PP_ALIGN.CENTER)
    add_footer(s, "Ablations: paper §6.1–6.3")

    # 12. Negative results and errors
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Negative results & error analysis", "失败实验说明了什么？剩余瓶颈在哪里？", 12)
    add_text(s, "没有奏效的方案", 0.74, 1.52, 5.55, 0.36, 20, RED, True)
    negs = [
        ("CRF", "0.6278 → 0.6249", "额外转移参数可能过拟合"),
        ("COMET 过筛", "0.3565 → 0.3418", "高阈值造成有效样本不足"),
        ("跨编码器", "0.3962 → 0.3721", "候选不一致引入噪声"),
        ("加权投票", "0.5008 → 0.4901", "置信度并非可靠质量代理"),
    ]
    for i, (name, delta, why) in enumerate(negs):
        y = 1.98 + i * 1.02
        add_rect(s, 0.74, y, 5.68, 0.82, WHITE, GRID, True, 1)
        add_text(s, name, 0.98, y + 0.15, 1.20, 0.28, 17.5, RED, True)
        add_text(s, delta, 2.24, y + 0.15, 1.70, 0.28, 16.5, INK, True, font=FONT_LATIN)
        add_text(s, why, 4.00, y + 0.15, 2.15, 0.40, 15.5, MUTED)
    add_text(s, "仍然存在的四类误差", 6.82, 1.52, 5.60, 0.36, 20, BLUE, True)
    errs = [
        ("边界错误", "日语与笔记本短语容易过扩张", BLUE),
        ("类别混淆", "实体-属性标签多，类别易混淆", PURPLE),
        ("低资源漏检", "俄/鞑/乌仍有较多漏检", ORANGE),
        ("V–A 收缩", "极端情绪回归向中心收缩", TEAL),
    ]
    for i, (name, body, c) in enumerate(errs):
        col, row = i % 2, i // 2
        x, y = 6.82 + col * 2.88, 2.00 + row * 1.65
        add_card(s, name, body, x, y, 2.65, 1.40, c, WHITE, 16.5, 14.5)
    add_rect(s, 6.82, 5.50, 5.70, 1.12, PALE_RED, "F3C4C9", True, 1)
    add_text(s, "局限：最终系统是 dataset-specific mixture；所有排名比较仍基于 dev；增强依赖近似精确的 span 对齐。", 7.05, 5.73, 5.25, 0.62, 17, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_footer(s, "Negative results, error analysis and limitations: paper §7–§8")

    # 13. Conclusion
    s = prs.slides.add_slide(blank)
    set_bg(s, NAVY)
    add_text(s, "CONCLUSION", 0.72, 0.45, 3.0, 0.28, 12, "9BC2E6", True, font=FONT_LATIN)
    add_text(s, "三句话带走 LEADER-DimASQP", 0.72, 0.92, 11.8, 0.58, 30, WHITE, True)
    cards = [
        ("01", "先解决结构稳定性", "Rethink + Hungarian matching 建立可靠的抽取与回归底座。", BLUE),
        ("02", "再按语言做选择", "分词、预训练、增强与阈值需要逐数据集验证，没有统一最优配方。", TEAL),
        ("03", "最后用集成降方差", "五种子集成平均 +4.76 pp，是最终系统最稳健的增益来源。", ORANGE),
    ]
    for i, (num, title, body, c) in enumerate(cards):
        x = 0.72 + i * 4.12
        add_rect(s, x, 2.02, 3.78, 2.88, "244A70", "3B6288", True, 1)
        add_text(s, num, x + 0.25, 2.28, 0.60, 0.36, 18, c, True, font=FONT_LATIN)
        add_text(s, title, x + 0.25, 2.84, 3.25, 0.48, 21, WHITE, True)
        add_text(s, body, x + 0.25, 3.54, 3.25, 0.92, 17.5, "DCE8F5")
    add_text(s, "0.5008 mean dev cF1  ·  +11.46 pp vs Llama-3.3 70B  ·  reported top on Jpn / Tat / Zho-R", 0.72, 5.45, 11.90, 0.45, 18.5, "70D7CD", True, PP_ALIGN.CENTER, font=FONT_LATIN)
    add_line(s, 2.52, 6.15, 10.82, 6.15, "3B6288", 1.2)
    add_text(s, "谢谢 · Questions?", 0.72, 6.47, 11.90, 0.50, 25, WHITE, True, PP_ALIGN.CENTER)

    # 14. References / backup
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, "Backup", "主要参考文献", 14)
    refs = [
        "[1] Lee et al. (2026). DimABSA: Building multilingual and multidomain datasets for dimensional ABSA. arXiv:2601.23022.",
        "[2] Yu et al. (2026). SemEval-2026 Task 3: Dimensional aspect-based sentiment analysis. arXiv:2604.07066.",
        "[3] Carion et al. (2020). End-to-end object detection with transformers. ECCV.",
        "[4] He et al. (2023). DeBERTaV3. ICLR.",
        "[5] Kuhn (1955). The Hungarian method for the assignment problem.",
        "[6] Sennrich et al. (2016). Improving NMT with monolingual data. ACL.",
        "[7] Rei et al. (2022). COMET-22 / COMETKiwi. WMT.",
    ]
    for i, ref in enumerate(refs):
        add_rect(s, 0.80, 1.54 + i * 0.72, 0.08, 0.34, BLUE if i < 2 else TEAL, BLUE if i < 2 else TEAL, False, 0)
        add_text(s, ref, 1.02, 1.50 + i * 0.72, 11.45, 0.52, 16.5, INK, i < 2, font=FONT_LATIN)
    add_text(s, "完整方法、消融、误差分析与局限请见 paper_iconip_current.tex。", 0.90, 6.83, 11.45, 0.34, 16, MUTED, True, PP_ALIGN.CENTER)
    add_footer(s)

    # Core metadata
    props = prs.core_properties
    props.title = "LEADER-DimASQP Paper Presentation (Revised)"
    props.subject = "Large-font, presentation-first paper talk"
    props.author = "Hongzhen Qin et al."
    props.comments = "Generated from docs/paper_iconip_current.tex and project experiment records."
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
