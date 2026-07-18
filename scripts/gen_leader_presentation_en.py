#!/usr/bin/env python3
"""Create the English LEADER-DimASQP talk from the revised master deck."""

from pathlib import Path
import re

from pptx import Presentation
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "LEADER_DimASQP_Presentation_revised.pptx"
OUTPUT = ROOT / "docs" / "LEADER_DimASQP_Presentation_English.pptx"


T = {
    # Cover
    "面向维度 ASQP 的语言特定编码器与数据感知集成系统": "A Language-Specific System for Dimensional ASQP",
    "核心结论": "CORE MESSAGE",
    "没有一个统一配方适合所有语言；稳定的集成与经验证的编码器选择决定最终表现。": "No universal recipe works best for every language: stable ensembling and validated encoder selection determine final performance.",
    "汇报人：秦洪振 · 铃木美美 · 福本文代 · 周仁杰": "Presenters: Hongzhen Qin · Yoshimi Suzuki · Fumiyo Fukumoto · Renjie Zhou",
    "杭州电子科技大学  ·  山梨大学": "Hangzhou Dianzi University  ·  University of Yamanashi",
    # Roadmap
    "用一条问题链讲清这篇论文": "A question-driven story of the paper",
    "任务是什么？": "What is the task?",
    "抽取结构 + 连续情绪回归": "Structure extraction + continuous affect regression",
    "为什么困难？": "Why is it hard?",
    "严格匹配、数据不平衡、回归收缩": "Exact matching, data imbalance, regression shrinkage",
    "系统怎么做？": "How does it work?",
    "Rethink、集合匹配、语言特定配置": "Rethink, set matching, per-language recipes",
    "证据是否充分？": "What is the evidence?",
    "基线、逐数据集结果与消融": "Baselines, per-dataset results, and ablations",
    "还能改什么？": "What remains?",
    "失败经验、误差与局限": "Negative results, errors, and limitations",
    "叙事主线：从任务约束出发，解释设计选择，再用实验验证，而不是按论文目录逐页复述。": "Narrative: start from task constraints, motivate each design choice, and then test it with evidence.",
    # Task
    "DimASQP：先抽对结构，才有资格评价情绪距离": "DimASQP: structure must match before affect distance counts",
    "输入句子（示意）": "Illustrative input",
    "输出集合：每条记录写成 (a, c, o, (V, A))": "Output set: each item is (a, c, o, (V, A))",
    "cF1 的两道门": "Two gates in cF1",
    "门 1 · categorical exact match": "Gate 1 · categorical exact match",
    "Aspect + Category + Opinion 必须匹配": "Aspect + Category + Opinion must match",
    "门 2 · V–A distance discount": "Gate 2 · V–A distance discount",
    "通过门 1 后，再按 V–A 欧氏距离折扣": "After Gate 1, credit is discounted by V–A distance",
    "因此：边界或类别错误会直接失去该条匹配；V–A 误差则连续扣分。": "Therefore, a span or category error loses the match; V–A error reduces credit continuously.",
    # Data
    "八个语言-领域数据集，规模与难点并不相同": "Eight language–domain datasets differ in size and bottlenecks",
    "训练句子数": "Training sentences",
    "中餐": "Zho-R",
    "英笔记本": "Eng-L",
    "中笔记本": "Zho-L",
    "英餐": "Eng-R",
    "日酒店": "Jpn-H",
    "俄餐": "Rus-R",
    "鞑餐": "Tat-R",
    "乌餐": "Ukr-R",
    "最大 / 最小 ≈ 4.9×": "Largest / smallest ≈ 4.9×",
    "挑战 1 · 严格结构匹配": "Challenge 1 · exact structural matching",
    "Aspect / Opinion 边界只偏 1–2 个词，也可能无法形成 categorical match。": "A 1–2 token shift in an Aspect or Opinion span can invalidate the categorical match.",
    "挑战 2 · 低资源与跨语差异": "Challenge 2 · low-resource variation",
    "训练规模从 1,240 到 6,050；分词、形态与领域标签分布各不相同。": "Training size ranges from 1,240 to 6,050; tokenization, morphology, and label distributions differ.",
    "挑战 3 · V–A 中心化": "Challenge 3 · V–A centralization",
    "点回归容易收缩到中性区域，极端情绪的区分度不足。": "Point regression tends to shrink toward neutral values, weakening separation at the extremes.",
    # Overview
    "LEADER-DimASQP = 稳定底座 + 按数据集选择": "LEADER-DimASQP = stable base + per-dataset selection",
    "A. 架构底座": "A. Architecture base",
    "先做 draft，再把结构与 V–A 软预测反馈给 token 表征；候选四元组用 Hungarian set matching 训练。": "Draft predictions → feed structural/V–A cues into token states → refine; train quad sets with Hungarian matching.",
    "B. 数据与迁移": "B. Data and transfer",
    "回译、paraphrase、COMET 过滤与跨语热启动都只在验证有效的数据集上保留。": "Keep BT, paraphrasing, COMET filtering, and warm starts only where dev results support them.",
    "C. 选择与推理": "C. Selection & inference",
    "按 dev cF1 选择编码器/阈值；五种子投票与 top-k 解码提升稳定性。": "Select encoders and thresholds by dev cF1; use five-seed voting and top-k decoding for stability.",
    # Architecture
    "两项结构设计：Rethink 反馈与集合级匹配": "Two structural designs: Rethink feedback and set-level matching",
    "BIO、类别、pair 与 V–A 草稿预测被嵌入": "Embed draft BIO, category, pair, and V–A predictions",
    "与原 token 表征融合后，由第二阶段预测头细化": "Fuse them with token states, then refine with a second-stage head",
    "在集合层面寻找一对一最优指派": "Find the optimal one-to-one assignment at set level",
    "减少候选顺序与重复配对带来的训练不稳定": "Reduce instability from candidate order and duplicate pairings",
    "训练稳定器：FGM ε=0.1  ·  EMA 0.999  ·  top-3 checkpoints": "Training stabilizers: FGM ε=0.1  ·  EMA 0.999  ·  top-3 checkpoints",
    "Feedback": "Cues",
    # Selection
    "最终系统不是一个模型，而是八个经验证的配置": "The final system is eight validated configurations—not one model",
    "强抽取底座": "Strong extraction base",
    "保留历史最优": "Retain validated best",
    "MeCab 裁剪；no para": "MeCab trimming; no para",
    "优于 ruRoberta": "Better than ruRoberta",
    "不保留 paraphrase": "No paraphrase retained",
    "五种子集成": "Five-seed ensemble",
    "原则：语言专用模型不一定更好；所有选择都必须由 dev cF1 验证。": "Principle: language-specific ≠ automatically better; validate every choice by dev cF1.",
    # Augmentation
    "增强不是越多越好：质量阈值与模型能力存在交互": "Augmentation quality interacts with model capacity",
    "日本语 paraphrase + COMET（mDeBERTa 阶段）": "Japanese paraphrase + COMET (mDeBERTa stage)",
    "结论 1 · 倒 U 型": "Finding 1 · inverted-U effect",
    "0.96 是该阶段的最佳阈值；提高到 0.965 后，样本量下降并出现欠拟合。": "At this stage, 0.96 is optimal; raising the threshold to 0.965 removes too much data and causes underfitting.",
    "结论 2 · 不能跨模型照搬": "Finding 2 · model-specific recipes",
    "换成 Japanese DeBERTa + 五种子集成后，no paraphrase 反而更好：0.4500 vs 0.4466。": "With Japanese DeBERTa and five-seed voting, no paraphrase is better: 0.4500 vs 0.4466.",
    "结论 3 · 迁移也需验证": "Finding 3 · validate transfer",
    "跨语热启动在早期低资源设置有效；最终俄/鞑使用 XLM-R-large，乌克兰保留 BT + paraphrase。": "Cross-lingual warm starts help early low-resource settings; the final Russian/Tatar systems use XLM-R-large, while Ukrainian keeps BT + paraphrase.",
    # Main results
    "平均 dev cF1 = 0.5008，显著超过官方基线": "Mean dev cF1 = 0.5008, well above the official baselines",
    # Per-dataset evidence
    "八项 dev 结果用于内部选择，不推导 test 排名": "Eight dev results guide internal selection—not test ranking",
    "论文可以自信地说": "What the paper can claim",
    "• 超过全部官方开发集基线\n• 平均 dev cF1 提升显著\n• 配方按数据集验证": "• Beats all official dev baselines\n• Improves mean dev cF1\n• Validates recipes per dataset",
    "论文不应说": "Unsupported claims",
    "• “所有数据集全面领先”\n• “绝对统治力”\n• 将 dev 结果写成 hidden-test 排名": "• ‘Best on every dataset’\n• ‘Absolute dominance’\n• Hidden-test rank from dev scores",
    "学术表达必须区分 dev 选择、official test 与外部系统。": "Keep dev selection, official test, and external systems distinct.",
    # Ablation
    "最稳健的最终贡献来自集成；其他收益依赖模型规模": "Ensembling is the most robust final gain; other effects depend on scale",
    "相对五种子单模均值": "versus the five-seed single-model mean",
    "8 个数据集全部提升": "Improves all eight datasets",
    "单项增益 +1.43 ～ +7.74 pp": "Per-dataset gain: +1.43 to +7.74 pp",
    "边界投票 + V–A 均值降低方差": "Boundary voting + V–A averaging reduce variance",
    "作用随 encoder/ensemble 变强而衰减": "The gain shrinks with stronger encoders and ensembles",
    "结论：language-specific ≠ automatically better": "Conclusion: language-specific ≠ automatically better",
    "论文主线应强调“验证驱动”，而不是把所有组件都包装成普适增益。": "The paper should emphasize validation-driven selection—not present every component as a universal gain.",
    # Negative results and errors
    "失败实验说明了什么？剩余瓶颈在哪里？": "What do failed experiments teach us, and what bottlenecks remain?",
    "没有奏效的方案": "What did not work",
    "额外转移参数可能过拟合": "Extra transition parameters may overfit",
    "COMET 过筛": "COMET filter",
    "高阈值造成有效样本不足": "A high threshold removes useful data",
    "跨编码器": "X-encoder fusion",
    "候选不一致引入噪声": "Inconsistent candidates introduce noise",
    "加权投票": "Weighted vote",
    "置信度并非可靠质量代理": "Confidence is not a reliable quality proxy",
    "仍然存在的四类误差": "Four remaining error types",
    "边界错误": "Span boundaries",
    "日语与笔记本短语容易过扩张": "Japanese and laptop spans often over-expand",
    "类别混淆": "Category confusion",
    "实体-属性标签多，类别易混淆": "Large entity–attribute inventories cause confusion",
    "低资源漏检": "Low-resource false negatives",
    "俄/鞑/乌仍有较多漏检": "Russian, Tatar, and Ukrainian still miss quads",
    "V–A 收缩": "V–A shrinkage",
    "极端情绪回归向中心收缩": "Extreme affect values regress toward the center",
    "局限：最终系统是 dataset-specific mixture；本 deck 仅报告 V20 dev 选择；增强依赖近似精确的 span 对齐。": "Limitations: the system is a dataset-specific mixture; this deck reports V20 dev selection only; augmentation relies on near-exact span alignment.",
    # Conclusion
    "三句话带走 LEADER-DimASQP": "LEADER-DimASQP in three takeaways",
    "先解决结构稳定性": "Stabilize structure first",
    "Rethink + Hungarian matching 建立可靠的抽取与回归底座。": "Rethink + Hungarian matching provide a reliable extraction and regression base.",
    "再按语言做选择": "Then select by language",
    "分词、预训练、增强与阈值需要逐数据集验证，没有统一最优配方。": "Tokenization, pretraining, augmentation, and thresholds require per-dataset validation.",
    "最后用集成降方差": "Finally, reduce variance",
    "五种子集成平均 +4.76 pp，是最终系统最稳健的增益来源。": "Five-seed ensembling adds +4.76 pp on average—the most robust final gain.",
    "谢谢 · Questions?": "Thank you · Questions?",
    # Backup
    "主要参考文献": "Key references",
    "完整方法、消融、误差分析与局限请见 paper_iconip_current.tex。": "See paper_iconip_current.tex for the full method, ablations, error analysis, and limitations.",
}


FONT_SIZES = {
    # Long slide titles
    "DimASQP: structure must match before affect distance counts": 22,
    "Eight language–domain datasets differ in size and bottlenecks": 22,
    "LEADER-DimASQP = stable base + per-dataset selection": 22,
    "Two structural designs: Rethink feedback and set-level matching": 20,
    "The final system is eight validated configurations—not one model": 20,
    "Augmentation quality interacts with model capacity": 22,
    "Mean dev cF1 = 0.5008, well above the official baselines": 22,
    "Gains are uneven: three reported tops, but four remaining gaps": 20,
    "Ensembling is the most robust final gain; other effects depend on scale": 20,
    "What do failed experiments teach us, and what bottlenecks remain?": 20,
    # Roadmap cards
    "What is the task?": 18,
    "Why is it hard?": 18,
    "How does it work?": 18,
    "What is the evidence?": 17,
    "What remains?": 18,
    "Structure extraction + continuous affect regression": 14,
    "Exact matching, data imbalance, regression shrinkage": 14,
    "Rethink, set matching, per-language recipes": 14,
    "Baselines, per-dataset results, and ablations": 14,
    "Negative results, errors, and limitations": 14,
    # Challenge cards
    "Challenge 1 · exact structural matching": 16.5,
    "Challenge 2 · low-resource variation": 16.5,
    "Challenge 3 · V–A centralization": 16.5,
    "A 1–2 token shift in an Aspect or Opinion span can invalidate the categorical match.": 15,
    "Training size ranges from 1,240 to 6,050; tokenization, morphology, and label distributions differ.": 15,
    "Point regression tends to shrink toward neutral values, weakening separation at the extremes.": 15,
    # Overview cards
    "A. Architecture base": 16.5,
    "B. Data and transfer": 16.5,
    "C. Selection & inference": 15.5,
    "Draft predictions → feed structural/V–A cues into token states → refine; train quad sets with Hungarian matching.": 13.5,
    "Keep BT, paraphrasing, COMET filtering, and warm starts only where dev results support them.": 13.5,
    "Select encoders and thresholds by dev cF1; use five-seed voting and top-k decoding for stability.": 13.5,
    # Augmentation cards
    "Finding 1 · inverted-U effect": 16.5,
    "Finding 2 · model-specific recipes": 16,
    "Finding 3 · validate transfer": 16,
    "At this stage, 0.96 is optimal; raising the threshold to 0.965 removes too much data and causes underfitting.": 14,
    "With Japanese DeBERTa and five-seed voting, no paraphrase is better: 0.4500 vs 0.4466.": 14,
    "Cross-lingual warm starts help early low-resource settings; the final Russian/Tatar systems use XLM-R-large, while Ukrainian keeps BT + paraphrase.": 13.5,
    # Per-dataset claim cards
    "What the paper can claim": 16.5,
    "Unsupported claims": 16.5,
    "• Beats all official dev baselines\n• Improves mean dev cF1\n• Validates recipes per dataset": 14.5,
    "• ‘Best on every dataset’\n• ‘Absolute dominance’\n• Hidden-test rank from dev scores": 14.5,
    "Keep dev selection, official test, and external systems distinct.": 13.5,
    # Ablation card copy
    "versus the five-seed single-model mean": 14.5,
    "Improves all eight datasets": 15,
    "Per-dataset gain: +1.43 to +7.74 pp": 15,
    "Boundary voting + V–A averaging reduce variance": 15,
    "The gain shrinks with stronger encoders and ensembles": 14,
    "The paper should emphasize validation-driven selection—not present every component as a universal gain.": 15,
    # Negative results and error cards
    "What did not work": 17,
    "Four remaining error types": 17,
    "COMET filter": 14,
    "X-encoder fusion": 14,
    "Weighted vote": 14,
    "Extra transition parameters may overfit": 12.5,
    "A high threshold removes useful data": 12.5,
    "Inconsistent candidates introduce noise": 12.5,
    "Confidence is not a reliable quality proxy": 12.5,
    "Span boundaries": 14.5,
    "Category confusion": 14.5,
    "Low-resource false negatives": 13.5,
    "V–A shrinkage": 14.5,
    "Japanese and laptop spans often over-expand": 12.5,
    "Large entity–attribute inventories cause confusion": 12.5,
    "Russian, Tatar, and Ukrainian still miss quads": 12.5,
    "Extreme affect values regress toward the center": 12.5,
    "Limitations: the system is a dataset-specific mixture; this deck reports V20 dev selection only; augmentation relies on near-exact span alignment.": 14.5,
    # Long validation principles
    "Principle: language-specific ≠ automatically better; validate every choice by dev cF1.": 15,
    # Compact labels and values in narrow boxes
    "Training sentences": 14,
    "6,050": 11.5,
    "4,076": 11.5,
    "3,490": 11.5,
    "2,284": 11.5,
    "1,600": 11.5,
    "1,240": 11.5,
    "a: food\nc: FOOD#QUALITY\no: excellent  ·  V–A: 8.2, 6.4": 12.5,
    "a: service\nc: SERVICE#GENERAL\no: cold  ·  V–A: 2.8, 4.1": 12.5,
}


def main():
    prs = Presentation(SOURCE)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text in T:
                        run.text = T[run.text]
                    run.font.name = "Aptos"
                    if run.text in FONT_SIZES:
                        run.font.size = Pt(FONT_SIZES[run.text])

    # Guard against accidentally shipping untranslated Chinese copy.
    leftovers = []
    cjk = re.compile(r"[\u3400-\u9fff]")
    for si, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and cjk.search(shape.text):
                leftovers.append((si, shape.text))
    if leftovers:
        raise RuntimeError(f"Untranslated text remains: {leftovers}")

    props = prs.core_properties
    props.title = "LEADER-DimASQP Paper Presentation (English)"
    props.subject = "English paper presentation with large-font projection layout"
    props.comments = "English version generated from the revised presentation master."
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
