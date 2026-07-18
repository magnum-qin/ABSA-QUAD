#!/usr/bin/env python3
"""Add a non-destructive superseded-evidence banner to historical decks."""

from argparse import ArgumentParser
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


def mark(path: Path, note: str) -> None:
    prs = Presentation(path)
    for slide in prs.slides:
        box = slide.shapes.add_textbox(
            Inches(0.08), Inches(0.02), prs.slide_width - Inches(0.16), Inches(0.30)
        )
        frame = box.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = Inches(0.04)
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = note
        run.font.name = "Aptos"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
    prs.save(path)


def main() -> None:
    parser = ArgumentParser()
    parser.add_argument("paths", nargs="+", type=Path)
    parser.add_argument(
        "--note",
        default=(
            "HISTORICAL DECK — superseded by the locked 2026-07-18 evidence; "
            "do not cite cross-split SOTA or universal KD claims"
        ),
    )
    args = parser.parse_args()
    for path in args.paths:
        mark(path, args.note)
        print(path)


if __name__ == "__main__":
    main()
